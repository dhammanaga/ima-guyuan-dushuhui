# -*- coding: utf-8 -*-
"""第5版 Word 生成器
用法: python3 mk_word5.py <系列交付目录>
把该目录下的 主持人版/学员版/00_主持人总指南/原文全集 md 转 docx 到 Word版/。
图片按 md 相对路径从系列目录内嵌；A4 2cm；页眉页脚页码。
"""
import os, sys, re, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HOST = (0x1F, 0x3A, 0x5F)   # 藏青
STU  = (0x1E, 0x46, 0x20)   # 墨绿
GRAY = (0x44, 0x4A, 0x55)

def shade(el, color):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), color)
    el.append(sh)

def left_bar(p, color):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6"); left.set(qn("w:color"), color)
    pbdr.append(left); pPr.append(pbdr); shade(pPr, "F2F7FB")

def add_runs(p, text, color=None):
    text = re.sub(r"`([^`]*)`", r"\1", text)
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg: continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg)
        if color: r.font.color.rgb = color

def add_field(p, field):
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = field
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)

def set_font(run, cn="宋体", size=None, color=None, bold=None):
    run.font.name = cn
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if size: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = color
    if bold is not None: run.bold = bold

def scale_image(path):
    """返回 (w,h) EMU，限制宽度 <= 16.5cm"""
    from PIL import Image
    try:
        with Image.open(path) as im:
            w, h = im.size
    except Exception:
        return None
    maxw = Cm(16.5)
    if w <= 0: return None
    ratio = min(1.0, maxw / Emu(int(w * 9525))) if False else None
    # px -> EMU: px*9525
    wemu = w * 9525; hemu = h * 9525
    if wemu > maxw:
        hemu = int(hemu * maxw / wemu); wemu = maxw
    return wemu, hemu

def convert(md_path, docx_path, color, header_text):
    txt = open(md_path, encoding="utf-8").read()
    base = os.path.dirname(os.path.abspath(md_path))
    # 预处理：把写在引用块内的表格行（> |...|）去掉引用符，令其成为普通表格
    lines = []
    raw = txt.split("\n")
    k = 0
    while k < len(raw):
        s = raw[k]
        m = re.match(r"^(?:>\s*)+(\|.*\|)\s*$", s)
        if m:
            # 收集连续的引用表格行
            grp = []
            while k < len(raw):
                mm = re.match(r"^(?:>\s*)+(\|.*\|)\s*$", raw[k])
                if not mm: break
                grp.append(mm.group(1)); k += 1
            lines.extend(grp)
        else:
            lines.append(s); k += 1
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    hp = sec.header.paragraphs[0]; hp.text = header_text
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs: set_font(r, "宋体", 8, RGBColor(*color))
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE"); fp.add_run(" / "); add_field(fp, "NUMPAGES")
    for r in fp.runs: set_font(r, "宋体", 8, RGBColor(*GRAY))
    st = doc.styles["Normal"]; st.font.name = "宋体"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = st.paragraph_format; pf.line_spacing = 1.3; pf.space_after = Pt(0)

    def H(text, lvl):
        p = doc.add_paragraph()
        if lvl == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(re.sub(r"（主持人版）|（学员版）", "", text))
            set_font(r, "微软雅黑", 18, RGBColor(*color), True)
        else:
            r = p.add_run(text)
            set_font(r, "微软雅黑", 15 if lvl == 1 else (13 if lvl == 2 else 12), RGBColor(*color), True)
            if lvl == 1: p.paragraph_format.space_before = Pt(10)
        return p

    def add_table(rows):
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=cols); t.style = "Table Grid"
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(cols):
                cell = cells[ci]; cell.text = ""
                p = cell.paragraphs[0]
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                val = row[ci] if ci < len(row) else ""
                for seg in re.split(r"(\*\*.+?\*\*)", val):
                    if not seg: continue
                    r = p.add_run(seg[2:-2] if seg.startswith("**") else seg)
                    if seg.startswith("**"): r.bold = True
                    set_font(r, "宋体", 9.5, RGBColor(0xFF,0xFF,0xFF) if ri==0 else None)
                if ri == 0:
                    shade(cell._tc.get_or_add_tcPr(), "%02X%02X%02X" % color)
                elif ri % 2 == 0:
                    shade(cell._tc.get_or_add_tcPr(), "F8FAFC")

    i = 0; n = len(lines)
    while i < n:
        s = lines[i].rstrip()
        # table
        if s.startswith("|") and i+1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i+1].strip()):
            rows = []; j = i
            while j < n and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")]); j += 1
            rows = [rows[0]] + rows[2:]
            add_table(rows); i = j; continue
        # image
        mi = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", s)
        if mi:
            alt, src = mi.group(1), mi.group(2)
            path = src if os.path.isabs(src) else os.path.join(base, src)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(path):
                wh = scale_image(path)
                r = p.add_run()
                if wh: r.add_picture(path, width=wh[0], height=wh[1])
                else: r.add_picture(path)
            else:
                r = p.add_run("［缺图：%s］" % alt); set_font(r, "宋体", 10, RGBColor(0xC0,0,0))
            if alt:
                cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = cp.add_run(alt); set_font(rr, "宋体", 9, RGBColor(*GRAY))
            i += 1; continue
        # heading
        mh = re.match(r"^(#{1,4}) (.*)$", s)
        if mh:
            H(mh.group(2).strip(), len(mh.group(1))); i += 1; continue
        # hr
        if re.match(r"^-{3,}$", s) or re.match(r"^\*{3,}$", s):
            p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"6"); bottom.set(qn("w:space"),"1"); bottom.set(qn("w:color"),"CCCCCC")
            pbdr.append(bottom); pPr.append(pbdr); i += 1; continue
        # list
        ml = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)$", s)
        if ml:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            marker = "• " if ml.group(2) in ("-", "*") else ""
            tail = ml.group(3)
            add_runs(p, marker + tail)
            i += 1; continue
        # blockquote
        if s.startswith(">"):
            buf = []
            while i < n and (lines[i].startswith(">") or lines[i].strip() == ">"):
                buf.append(re.sub(r"^(?:>\s*)+", "", lines[i]).strip()); i += 1
            for b in buf:
                if not b: continue
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3)
                add_runs(p, b); left_bar(p, "%02X%02X%02X" % color)
            continue
        if s.startswith("<div") or s.startswith("</div") or s.startswith("<!--"):
            i += 1; continue
        if not s:
            i += 1; continue
        p = doc.add_paragraph(); add_runs(p, s); i += 1
    doc.save(docx_path)

def main(folder):
    outdir = os.path.join(folder, "Word版"); os.makedirs(outdir, exist_ok=True)
    for md in sorted(glob.glob(os.path.join(folder, "*.md"))):
        name = os.path.splitext(os.path.basename(md))[0]
        if "学员版" in name: color = STU
        elif "主持人版" in name: color = HOST
        else: color = HOST
        # 页眉
        mw = re.search(r"(第\d+周[^_]*?)_", name)
        if "第" in name and "周" in name:
            wk = name.split("_")[0]
            role = "学员版" if "学员版" in name else ("主持人版" if "主持人版" in name else name)
            htt = f"{wk}　|　{role}"
        else:
            htt = name
        convert(md, os.path.join(outdir, name + ".docx"), color, htt)
        print("生成", name + ".docx")
    print("完成 ->", outdir)

if __name__ == "__main__":
    main(sys.argv[1])
