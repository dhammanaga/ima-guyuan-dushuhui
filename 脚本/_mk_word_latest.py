#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
_mk_word_latest.py —— 读书会资料包 md → docx 生成器（按《读书会 Word 文档生成规则·最新版 v1.0》实现）

规则要点（与 2026-09-28 第四轮 Word 产物实测一致）：
  页面 A4 21x29.7cm，四边距 2cm；Normal 西文等线 + 中文宋体 10.5pt，单倍行距
  主持人版主题色 #1F3864（藏青），学员版 #1E4620（墨绿）
  首行 # 标题 18pt 加粗居中（去「（主持人版）/（学员版）」后缀）；# 15pt；## 13pt；###/#### 11.5pt
  页眉 8pt 主题色居中：读书会NN · 第N周 · {edition}（4周版|8周版） · 主持人版|学员版
  页脚 居中 PAGE 域
  引用块 每行独立成段：左缩进 0.3cm + 左侧主题色竖线 sz18 + 底纹 F2F5FA
  表格 Table Grid：表头主题色底白字 9pt，数据行隔行 F2F5FA，单元格 9pt
  列表 - → List Bullet；1. → List Number（去原编号）；--- / *** / HTML 注释丢弃
  时间戳 <div align="right">…</div> → 右对齐 8pt #888888
  不嵌入图片

用法：
  python3 _mk_word_latest.py <资料包目录|单个md> [--series 05] [--edition 第八轮去禅修引导第8版]
                            [--weeks 4|8] [--out 输出目录] [--zip] [--dry]
"""
import os, re, sys, glob, zipfile, argparse, datetime, collections

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HOST_C = RGBColor(0x1F, 0x38, 0x64)   # 主持人版 藏青
STU_C = RGBColor(0x1E, 0x46, 0x20)    # 学员版 墨绿
QUOTE_BG = "F2F5FA"
STAMP_GREY = RGBColor(0x88, 0x88, 0x88)
SKIP_NAMES = ("README", "更新日志_CHANGELOG", "CHANGELOG")

WEEK_RE = re.compile(r"^第(\d+)周_(.+)_(主持人版|学员版)\.md$")


# ---------- 底层工具 ----------
def _shade(el, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def _left_bar(p, color_hex):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), color_hex)
    pbdr.append(left)
    pPr.append(pbdr)
    _shade(pPr, QUOTE_BG)


def _field(paragraph, code):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)


def add_runs(p, text, color=None, size=None, base_bold=False):
    """处理段内 **加粗** 与 `代码`，保留其余标记文字。"""
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = text.replace("<br>", " ").replace("<br/>", " ")
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg); r.bold = base_bold
        if color is not None:
            r.font.color.rgb = color
        if size is not None:
            r.font.size = Pt(size)


# ---------- 单文件转换 ----------
def convert(md_path, docx_path, header_text, color, color_hex):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)

    hp = sec.header.paragraphs[0]
    hp.text = header_text
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.size = Pt(8); r.font.color.rgb = color

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, "PAGE")
    for r in fp.runs:
        r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    st = doc.styles["Normal"]
    st.font.name = "等线"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def heading(text, lvl, first=False):
        p = doc.add_paragraph()
        if first:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t = text.replace("（主持人版）", "").replace("（学员版）", "").strip()
            r = p.add_run(t); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = color
            return
        p.add_run(text).bold = True
        p.runs[0].font.color.rgb = color
        p.runs[0].font.size = Pt(15 if lvl == 1 else (13 if lvl == 2 else 11.5))
        if lvl == 1:
            p.paragraph_format.space_before = Pt(10)

    def make_table(rows):
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(cols):
                p = cells[ci].paragraphs[0]
                add_runs(p, row[ci] if ci < len(row) else "",
                         color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
                for r in p.runs:
                    r.font.size = Pt(9)
                tcPr = cells[ci]._tc.get_or_add_tcPr()
                if ri == 0:
                    _shade(tcPr, color_hex)
                elif ri % 2 == 0:
                    _shade(tcPr, QUOTE_BG)

    def dequote(line):
        """剥掉行首连续的 > 前缀（用于识别引用块内嵌表格）"""
        s = line.rstrip()
        while s.lstrip().startswith(">"):
            s = s.lstrip()[1:]
        return s.strip()

    i, n, seen_content = 0, len(lines), False
    while i < n:
        s = lines[i].rstrip()
        stripped = s.strip()

        # 表格块（含引用块内嵌表格：> | a | b |）
        cur = dequote(lines[i])
        nxt = dequote(lines[i + 1]) if i + 1 < n else ""
        if cur.startswith("|") and re.match(r"^\|[\s:|\-]+\|$", nxt):
            rows, j = [], i
            while j < n:
                row = dequote(lines[j])
                if not row.startswith("|"):
                    break
                rows.append([c.strip() for c in row.strip("|").split("|")])
                j += 1
            if len(rows) >= 2:
                rows = [rows[0]] + rows[2:]
            make_table(rows)
            i = j; seen_content = True; continue

        # 标题
        m = re.match(r"^(#{1,6}) (.*)$", s)
        if m:
            lvl, text = len(m.group(1)), m.group(2).strip()
            heading(text, lvl, first=(not seen_content))
            i += 1; seen_content = True; continue

        # 引用块
        if stripped.startswith(">"):
            buf = []
            while i < n and (lines[i].lstrip().startswith(">") or lines[i].strip() == ">"):
                # 引用块内嵌表格：停止收集，交回外层表格分支
                cc = dequote(lines[i])
                nn2 = dequote(lines[i + 1]) if i + 1 < n else ""
                if cc.startswith("|") and re.match(r"^\|[\s:|\-]+\|$", nn2):
                    break
                buf.append(lines[i].lstrip()[1:].strip()); i += 1
            if not buf and i < n:
                # 防御：当前行即引用块内表格起点而未前进时，强制前进来避免死循环
                i += 1
            for b in buf:
                if not b or b.startswith("<!--"):
                    continue
                p = doc.add_paragraph()
                add_runs(p, b)
                p.paragraph_format.left_indent = Cm(0.3)
                _left_bar(p, color_hex)
            seen_content = True; continue

        # 时间戳 div
        if stripped.startswith("<div") or stripped.startswith("</div>"):
            mm = re.search(r">(.*?)</div>", s)
            if mm and mm.group(1).strip():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(mm.group(1).strip())
                r.font.size = Pt(8); r.font.color.rgb = STAMP_GREY
            i += 1; continue

        # HTML 注释 / 分隔线 / 空行
        if stripped.startswith("<!--") or stripped in ("---", "***", "___"):
            i += 1; continue
        if not stripped:
            i += 1; continue

        # 列表
        if re.match(r"^[-*] ", s):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, s[2:]); i += 1; seen_content = True; continue
        if re.match(r"^\d+\. ", s):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s*", "", s)); i += 1; seen_content = True; continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1; seen_content = True

    os.makedirs(os.path.dirname(docx_path) or ".", exist_ok=True)
    doc.save(docx_path)
    return doc


# ---------- 批量 ----------
def infer_series(path):
    m = re.search(r"读书会(\d{1,2})", path)
    return m.group(1).zfill(2) if m else "00"


def infer_weeks(path):
    return 8 if "8周" in path else 4


def build_header(nn, md_name, edition, weeks):
    tag = f"（{weeks}周版）"
    m = WEEK_RE.match(md_name)
    if m:
        wk, _, role = m.group(1), m.group(2), m.group(3)
        return f"读书会{nn} · 第{wk}周 · {edition}{tag} · {role}"
    stem = os.path.splitext(md_name)[0]
    if stem.startswith("00_"):
        return f"读书会{nn} · 主持人总指南 · {edition}{tag}"
    return f"读书会{nn} · {stem} · {edition}{tag}"


def convert_package(src, out=None, series=None, edition="第八轮（第8版）", weeks=None, do_zip=False):
    src = os.path.abspath(src)
    pkg = src if os.path.isdir(src) else os.path.dirname(src)
    outdir = out or os.path.join(pkg, "Word版")
    os.makedirs(outdir, exist_ok=True)

    mds = ([src] if os.path.isfile(src)
           else sorted(glob.glob(os.path.join(pkg, "*.md"))))
    targets = [f for f in mds
               if not os.path.basename(f).startswith(SKIP_NAMES)]

    stats = collections.Counter()
    for md in targets:
        bn = os.path.basename(md)
        m = WEEK_RE.match(bn)
        stu = bool(m and m.group(3) == "学员版")
        color, hexc = (STU_C, "1E4620") if stu else (HOST_C, "1F3864")
        nn = series or infer_series(pkg)
        wks = weeks or infer_weeks(pkg)
        header = build_header(nn, bn, edition, wks)
        docx_path = os.path.join(outdir, bn.replace(".md", ".docx"))
        doc = convert(md, docx_path, header, color, hexc)
        stats["docx"] += 1
        stats["tables"] += len(doc.tables)
        stats["quotes"] += sum(1 for p in doc.paragraphs if "w:pBdr" in p._p.xml)
        sizes = collections.Counter(
            p.runs[0].font.size.pt for p in doc.paragraphs
            if p.runs and p.runs[0].font.size)
        print(f"  ✅ {bn[:-3]} → {os.path.basename(docx_path)} | 表 {len(doc.tables)} | "
              f"引用段 {sum(1 for p in doc.paragraphs if 'w:pBdr' in p._p.xml)} | 字号 {dict(sorted(sizes.items()))}")

    if do_zip:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M")
        zpath = os.path.join(pkg, f"{os.path.basename(pkg)}_Word版_{ts}_v1.0.zip")
        with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(glob.glob(os.path.join(outdir, "*.docx"))):
                z.write(f, os.path.join("Word版", os.path.basename(f)))
        print(f"  📦 {os.path.basename(zpath)}（{os.path.getsize(zpath)/1024:.0f} KB）")

    print(f"合计：docx {stats['docx']} 份 | 表格 {stats['tables']} | 引用段 {stats['quotes']} | 输出 {outdir}")
    return stats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("--series")
    ap.add_argument("--edition", default="第八轮（第8版）")
    ap.add_argument("--weeks", type=int, choices=(4, 8))
    ap.add_argument("--out")
    ap.add_argument("--zip", action="store_true")
    ap.add_argument("--dry", action="store_true")
    a = ap.parse_args()
    if a.dry:
        print("dry-run：", a.src, a.series, a.edition, a.weeks)
        return
    convert_package(a.src, a.out, a.series, a.edition, a.weeks, a.zip)


if __name__ == "__main__":
    main()
