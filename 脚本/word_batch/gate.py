#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""逐份门禁：5 项硬检查（任一不过即 FAIL）"""
import os, re, sys, json, glob, zipfile, importlib.util, io, contextlib

sys.path.insert(0, "/sandbox/workspace/scripts")
import verify_word as VW
import verify_md_docx as VM
from docx import Document

EDITION = "第八轮（第8版）"
WEEKS = 4
WHITELIST = {"宋体", "Times New Roman", "微软雅黑"}


def fonts_of(path):
    z = zipfile.ZipFile(path)
    fonts = set()
    for n in z.namelist():
        if not n.endswith(".xml"):
            continue
        x = z.read(n).decode("utf-8", "ignore")
        for m in re.finditer(r"<w:rFonts\b[^>]*>", x):
            for a in re.findall(r'w:(?:ascii|hAnsi|eastAsia|cs)="([^"]*)"', m.group(0)):
                if a:
                    fonts.add(a)
        if n == "word/theme/theme1.xml":
            for tag in ("majorFont", "minorFont"):
                blk = re.search(r"<a:%s>(.*?)</a:%s>" % (tag, tag), x, re.S)
                if blk:
                    for a in re.findall(r'<a:(?:latin|ea|cs)\s+typeface="([^"]*)"', blk.group(1)):
                        if a:
                            fonts.add(a)
    return fonts


def resid_of(path):
    """md 语法残留判定：
       ** / <!-- / ```  → 任何出现即残留
       #                → 仅「行首 #{1,6} 空格」属残留（正文讲规则时的字面 ### 属源文内容）
       |                → 仅「行首 | 」或「≥2 个 |」（未转表格的行）属残留；时间戳行单个 | 属源文内容
       另返回字符级明细供人工核。"""
    d = Document(path)
    texts = [p.text for p in d.paragraphs]
    texts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    gate = {"**": 0, "<!--": 0, "```": 0, "#行首": 0, "|表格行": 0}
    chars = {"**": 0, "#": 0, "|": 0}
    for t in texts:
        chars["**"] += t.count("**")
        chars["#"] += t.count("#")
        chars["|"] += t.count("|")
        gate["**"] += t.count("**")
        gate["<!--"] += t.count("<!--")
        gate["```"] += t.count("```")
        if re.match(r"^\s*#{1,6}\s", t):
            gate["#行首"] += 1
        if t.strip().startswith("|") or t.count("|") >= 2:
            gate["|表格行"] += 1
    return gate, chars


def gate_one(rec):
    md, dx, nn = rec["md"], rec["docx"], rec["nn"]
    # 1 verify_word 12 项
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        ok1 = VW.check(dx, nn, WEEKS, EDITION)
    fails1 = [l.strip() for l in buf.getvalue().split("\n") if "[FAIL]" in l]
    # 2 md/docx 差异
    mdu, dxu = VM.parse_md(md), VM.parse_docx(dx)
    diffs = VM.compare(mdu, dxu)
    ex = VM.doc_extra_checks(dx)
    bad_fonts_vm = ex["bad_fonts"]
    # 3 numPr
    numpr = ex["numpr"]
    # 4 字体白名单
    fonts = fonts_of(dx)
    bad_fonts = sorted(f for f in fonts if f not in WHITELIST)
    # 5 md 残留
    hits, chars = resid_of(dx)
    ok = (ok1 and not diffs and numpr == 0 and not bad_fonts and
          sum(hits.values()) == 0 and not bad_fonts_vm)
    return dict(nn=nn, week=rec["week"], role=rec["role"], name=rec["name"],
                docx=dx, md=md, unit_md=len(mdu), unit_docx=len(dxu),
                vw="PASS" if ok1 else "FAIL", vw_fail=fails1,
                diff=len(diffs), diff_samples=diffs[:5], numpr=numpr,
                fonts=sorted(fonts), bad_fonts=bad_fonts, bad_fonts_vm=bad_fonts_vm,
                resid=hits, resid_chars=chars, gate="PASS" if ok else "FAIL")


NAME_RE = re.compile(r"^读书会(\d{2})_(第\d+周_.+_(?:主持人版|学员版))\.docx$")


def collect():
    recs = []
    for d in sorted(glob.glob("/sandbox/workspace/out/读书会*")):
        nn = os.path.basename(d)[-2:]
        for f in sorted(glob.glob(os.path.join(d, "*.docx"))):
            bn = os.path.basename(f)
            m = NAME_RE.match(bn)
            assert m, bn
            md = os.path.join("/sandbox/workspace/src", "读书会" + nn, m.group(2) + ".md")
            assert os.path.exists(md), md
            recs.append(dict(nn=nn, week=int(re.search(r"第(\d+)周", m.group(2)).group(1)),
                             role="学员版" if bn.endswith("学员版.docx") else "主持人版",
                             md=md, docx=f, name=bn))
    return recs


def main(names):
    stats = collect()
    todo = [r for r in stats if r["name"] in names] if names else stats
    out = []
    for r in todo:
        g = gate_one(r)
        out.append(g)
        print(("[PASS] " if g["gate"] == "PASS" else "[FAIL] ") + g["name"] +
              f" | diff={g['diff']} numPr={g['numpr']} 残留={sum(g['resid'].values())} 字符级={g['resid_chars']}"
              + ("" if g["gate"] == "PASS" else f" | vw_fail={g['vw_fail']} bad_fonts={g['bad_fonts']} resid={g['resid']} diff={g['diff_samples']}"))
        sys.stdout.flush()
    p = "/sandbox/workspace/work/gate_%s.json" % (names[0] if names else "all")
    json.dump(out, open(p, "w"), ensure_ascii=False, indent=1)
    bad = [g for g in out if g["gate"] != "PASS"]
    print(f"\n合计 {len(out)} 份 | PASS {len(out)-len(bad)} | FAIL {len(bad)}")
    print("FAIL 清单:", [g["name"] for g in bad])


if __name__ == "__main__":
    main(sys.argv[1:])
