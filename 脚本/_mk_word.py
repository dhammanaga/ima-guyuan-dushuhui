#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
_mk_word.py <NN> —— 第三轮4周版 Word 生成（D6）
把 4周版各周主持人版/学员版 + 00_主持人总指南.md 转 docx（主持人版藏青、学员版墨绿；A4 2cm；页眉/页脚）。
用法: python3 _mk_word.py <NN>
"""
import os, sys, re, glob, json, zipfile, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/sandbox/workspace/新型读书会.大佛寺/第二轮"
OUT_ROOT = os.path.join(BASE, "交付成果/第三轮修订版/第三轮4周版")
CFG = json.load(open(os.path.join(BASE, "_thirdround_titles.json"), encoding="utf-8"))
HOST_C = RGBColor(0x1F, 0x38, 0x64)   # 藏青
STU_C = RGBColor(0x1E, 0x46, 0x20)    # 墨绿

def shade(el, color):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), color)
    el.append(sh)

def set_left_bar(p, color):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18"); left.set(qn("w:space"), "4"); left.set(qn("w:color"), color)
    pbdr.append(left); pPr.append(pbdr); shade(pPr, "F2F5FA")

def add_runs(p, text, color=None, base_bold=False):
    text = re.sub(r"`([^`]*)`", r"\1", text)
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg); r.bold = base_bold
        if color:
            r.font.color.rgb = color

def add_field(paragraph, field):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = field
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)

def convert(md_path, docx_path, title, header_text, color, student=False):
    txt = open(md_path, encoding="utf-8").read()
    lines = txt.split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    # 页眉
    hp = sec.header.paragraphs[0]; hp.text = header_text
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.size = Pt(8); r.font.color.rgb = color
    # 页脚页码
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")
    st = doc.styles["Normal"]; st.font.name = "等线"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def H(text, lvl):
        p = doc.add_paragraph()
        if lvl == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text.replace("（主持人版）", "").replace("（学员版）", ""))
            r.bold = True; r.font.size = Pt(18); r.font.color.rgb = color
            return
        r = p.add_run(text)
        r.bold = True; r.font.color.rgb = color
        r.font.size = Pt(15 if lvl == 1 else (13 if lvl == 2 else 11.5))
        if lvl == 1:
            p.paragraph_format.space_before = Pt(10)
        return

    def table(rows):
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=cols); t.style = "Table Grid"
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(cols):
                cell = cells[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                add_runs(p, row[ci] if ci < len(row) else "", color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
                for r in p.runs:
                    r.font.size = Pt(9)
                if ri == 0:
                    shade(cell._tc.get_or_add_tcPr(), "%02X%02X%02X" % (color[0] if False else 0x1F, 0x38, 0x64) if not student else "1E4620")
                elif ri % 2 == 0:
                    shade(cell._tc.get_or_add_tcPr(), "F2F5FA")
        return t

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.rstrip()
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            rows = [rows[0]] + rows[2:]
            table(rows); i = j; continue
        if re.match(r"^#{1,4} ", s):
            lvl = len(s) - len(s.lstrip("#"))
            H(s[lvl + 1:].strip(), lvl); i += 1; continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and (lines[i].startswith(">") or lines[i].strip() == ">"):
                buf.append(lines[i].lstrip(">").strip()); i += 1
            for b in buf:
                if not b:
                    continue
                p = doc.add_paragraph(); add_runs(p, b)
                p.paragraph_format.left_indent = Cm(0.3)
                set_left_bar(p, "%02X%02X%02X" % (color[0], color[1], color[2]))
            continue
        if s.startswith("<div") or s.startswith("</div>"):
            m = re.search(r">(.*?)</div>", s)
            if m:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(m.group(1)); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            i += 1; continue
        if re.match(r"^[-*] ", s):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, s[2:]); i += 1; continue
        if re.match(r"^\d+\. ", s):
            p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\.\s*", "", s)); i += 1; continue
        if s.strip() in ("---", "***"):
            i += 1; continue
        if not s.strip():
            i += 1; continue
        p = doc.add_paragraph(); add_runs(p, s); i += 1
    doc.save(docx_path)

def run(NN):
    cfg = CFG[str(NN)]
    pkg = glob.glob(os.path.join(OUT_ROOT, f"读书会{NN:02d}_*_4周资料包（第三版修订）"))[0]
    wdir = os.path.join(pkg, "Word版"); os.makedirs(wdir, exist_ok=True)
    for md in sorted(glob.glob(pkg + "/*.md")):
        bn = os.path.basename(md)
        if bn in ("README_资料包说明.md", "更新日志_CHANGELOG.md", "原文全集.md"):
            continue
        m = re.match(r"第(\d+)周_(.*)_(主持人版|学员版)\.md$", bn)
        if m:
            wk, tt, ver = m.group(1), m.group(2), m.group(3)
            stu = ver == "学员版"
            convert(md, os.path.join(wdir, bn.replace(".md", ".docx")),
                    tt, f"读书会{NN:02d} · 第{wk}周 · 第三版修订（4周版） · {ver}", STU_C if stu else HOST_C, stu)
        elif bn.startswith("00_"):
            convert(md, os.path.join(wdir, "00_主持人总指南.docx"), "主持人总指南",
                    f"读书会{NN:02d} · 主持人总指南 · 第三版修订（4周版）", HOST_C)
    # Word 包
    zpath = os.path.join(pkg, f"读书会{NN:02d}_4周资料包_Word版_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(glob.glob(wdir + "/*.docx")):
            z.write(f, os.path.join("Word版", os.path.basename(f)))
    print(f"✅ Word {NN:02d}: {len(glob.glob(wdir+'/*.docx'))} docx, {os.path.getsize(zpath)/1024:.0f} KB")

if __name__ == "__main__":
    run(int(sys.argv[1]))
